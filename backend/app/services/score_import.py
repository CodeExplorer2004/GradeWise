from __future__ import annotations

import csv
import io
import json
import unicodedata
from collections.abc import Iterable
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.models import (
    ClassMembership,
    Exam,
    ExamScore,
    GradeSubjectConfig,
    ImportBatch,
    SchoolClass,
    Student,
    Subject,
    User,
)
from app.core.schemas import ImportErrorItem, ImportResult
from app.core.scoring import has_valid_score_increment

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2_000
SUPPORTED_SUFFIXES = {".csv": "csv", ".xlsx": "xlsx", ".json": "json"}
HEADER_ALIASES = {
    "student_no": "student_no",
    "学号": "student_no",
    "exam_name": "exam_name",
    "考试": "exam_name",
    "考试名称": "exam_name",
    "subject_code": "subject_code",
    "subject": "subject_code",
    "科目": "subject_code",
    "科目编码": "subject_code",
    "score": "score",
    "成绩": "score",
    "分数": "score",
}
REQUIRED_FIELDS = {"student_no", "exam_name", "subject_code", "score"}


def score_validation_error(score: float, max_score: float) -> str | None:
    if not 0 <= score <= max_score:
        return f"成绩必须在 0 到 {max_score:g} 之间"
    if not has_valid_score_increment(score):
        return "成绩只能为整数或 0.5 分"
    return None


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _canonical_row(row: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for key, value in row.items():
        canonical = HEADER_ALIASES.get(_clean_text(key).lower())
        if canonical:
            result[canonical] = _clean_text(value)
    return result


def _csv_rows(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    return [dict(row) for row in csv.DictReader(io.StringIO(text))]


def _xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [_clean_text(value) for value in next(rows, [])]
    return [dict(zip(headers, values, strict=False)) for values in rows]


def _json_rows(content: bytes) -> list[dict[str, Any]]:
    payload = json.loads(content.decode("utf-8-sig"))
    if isinstance(payload, dict):
        payload = payload.get("rows")
    if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
        raise ValueError("JSON 必须是对象数组，或包含 rows 数组")
    return payload


def parse_score_file(filename: str, content: bytes) -> tuple[str, list[dict[str, str]]]:
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("文件不能超过 5MB")
    suffix = Path(filename).suffix.lower()
    source_type = SUPPORTED_SUFFIXES.get(suffix)
    if not source_type:
        raise ValueError("仅支持 CSV、XLSX 和 JSON 文件")
    try:
        raw_rows = {
            "csv": _csv_rows,
            "xlsx": _xlsx_rows,
            "json": _json_rows,
        }[source_type](content)
    except (
        UnicodeDecodeError,
        csv.Error,
        json.JSONDecodeError,
        BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
    ) as exc:
        raise ValueError(f"文件解析失败：{exc}") from exc
    if not raw_rows:
        raise ValueError("文件中没有可导入的数据行")
    if len(raw_rows) > MAX_ROWS:
        raise ValueError(f"单次最多导入 {MAX_ROWS} 行")
    return source_type, [_canonical_row(row) for row in raw_rows]


def _unique_by(items: Iterable[Any], key: str) -> tuple[dict[str, Any], set[str]]:
    result: dict[str, Any] = {}
    duplicates: set[str] = set()
    for item in items:
        value = _clean_text(getattr(item, key))
        if value in result:
            duplicates.add(value)
        else:
            result[value] = item
    return result, duplicates


async def import_scores(
    session: AsyncSession,
    user: User,
    filename: str,
    content: bytes,
) -> ImportResult:
    safe_filename = (filename or "scores.csv").replace("\\", "/").rsplit("/", 1)[-1][:255]
    source_type, rows = parse_score_file(safe_filename, content)
    students = list(
        (await session.scalars(select(Student).where(Student.school_id == user.school_id))).all()
    )
    exams = list(
        (await session.scalars(select(Exam).where(Exam.school_id == user.school_id))).all()
    )
    subjects = list(
        (await session.scalars(select(Subject).where(Subject.school_id == user.school_id))).all()
    )
    memberships = list((await session.scalars(select(ClassMembership))).all())
    classes = list(
        (
            await session.scalars(
                select(SchoolClass).where(SchoolClass.school_id == user.school_id)
            )
        ).all()
    )
    curriculum = list(
        (
            await session.scalars(
                select(GradeSubjectConfig).where(
                    GradeSubjectConfig.school_id == user.school_id
                )
            )
        ).all()
    )
    existing_scores = list(
        (
            await session.scalars(
                select(ExamScore).join(Student).where(Student.school_id == user.school_id)
            )
        ).all()
    )

    student_map = {_clean_text(item.student_no): item for item in students}
    exam_map, ambiguous_exams = _unique_by(exams, "name")
    subject_map = {
        key: item
        for item in subjects
        for key in (_clean_text(item.code), _clean_text(item.name))
    }
    classes_by_id = {item.id: item for item in classes}
    membership_map = {
        (item.student_id, item.academic_year): classes_by_id.get(item.class_id)
        for item in memberships
    }
    curriculum_map = {
        (item.academic_year, item.grade_level, item.subject_id): item
        for item in curriculum
    }
    existing_map = {
        (item.student_id, item.exam_id, item.subject_id): item for item in existing_scores
    }

    errors: list[ImportErrorItem] = []
    seen: set[tuple[int, int, int]] = set()
    accepted = 0
    inserted = 0
    updated = 0
    for row_number, row in enumerate(rows, start=2):
        missing = REQUIRED_FIELDS - set(row)
        if missing or any(not row.get(field) for field in REQUIRED_FIELDS):
            errors.append(ImportErrorItem(row=row_number, reason="缺少学号、考试、科目或成绩"))
            continue
        student = student_map.get(row["student_no"])
        if not student:
            errors.append(ImportErrorItem(row=row_number, reason="学号不存在"))
            continue
        if row["exam_name"] in ambiguous_exams:
            errors.append(ImportErrorItem(row=row_number, reason="考试名称不唯一"))
            continue
        exam = exam_map.get(row["exam_name"])
        if not exam:
            errors.append(ImportErrorItem(row=row_number, reason="考试不存在"))
            continue
        subject = subject_map.get(row["subject_code"])
        if not subject:
            errors.append(ImportErrorItem(row=row_number, reason="科目编码或名称不存在"))
            continue
        try:
            score = float(row["score"])
        except ValueError:
            errors.append(ImportErrorItem(row=row_number, reason="成绩不是有效数字"))
            continue
        school_class = membership_map.get((student.id, exam.academic_year))
        if not school_class:
            errors.append(ImportErrorItem(row=row_number, reason="学生在该学年没有班级关系"))
            continue
        if exam.grade_level and exam.grade_level != school_class.grade_level:
            errors.append(ImportErrorItem(row=row_number, reason="考试年级与学生所在年级不一致"))
            continue
        subject_config = curriculum_map.get(
            (exam.academic_year, school_class.grade_level, subject.id)
        )
        if not subject_config:
            errors.append(ImportErrorItem(row=row_number, reason="该学年年级未开设此科目"))
            continue
        validation_error = score_validation_error(score, subject_config.max_score)
        if validation_error:
            errors.append(ImportErrorItem(row=row_number, reason=validation_error))
            continue
        score_key = (student.id, exam.id, subject.id)
        if score_key in seen:
            errors.append(ImportErrorItem(row=row_number, reason="文件内存在重复成绩记录"))
            continue
        seen.add(score_key)
        existing = existing_map.get(score_key)
        if existing:
            existing.score = score
            updated += 1
        else:
            session.add(
                ExamScore(
                    student_id=student.id,
                    class_id=school_class.id,
                    exam_id=exam.id,
                    subject_id=subject.id,
                    score=score,
                )
            )
            inserted += 1
        accepted += 1

    batch_id = str(uuid4())
    session.add(
        ImportBatch(
            id=batch_id,
            school_id=user.school_id,
            user_id=user.id,
            filename=safe_filename,
            source_type=source_type,
            status="completed" if not errors else "partial",
            total_rows=len(rows),
            accepted_rows=accepted,
            rejected_rows=len(errors),
            errors=[item.model_dump() for item in errors[:100]],
        )
    )
    await session.commit()
    return ImportResult(
        batch_id=batch_id,
        filename=safe_filename,
        source_type=source_type,
        total_rows=len(rows),
        accepted_rows=accepted,
        rejected_rows=len(errors),
        inserted_rows=inserted,
        updated_rows=updated,
        errors=errors[:100],
    )
